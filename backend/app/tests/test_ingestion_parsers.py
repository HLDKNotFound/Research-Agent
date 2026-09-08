import io
import math
import uuid
import docx
import fitz
import pandas as pd
import pytest
from app.ingestion.chunker import SemanticChunker
from app.ingestion.embeddings import embedding_engine
from app.ingestion.parsers import parse_document, parse_docx, parse_pdf, parse_spreadsheet
from app.models.project import Project, ProjectMember
from app.models.user import User
from app.schemas.file import FileCreateDTO
from app.services.file_service import FileService
from app.services.ingestion_service import IngestionService


@pytest.mark.asyncio
async def test_document_parsers_and_chunker():
    # 1. Test CSV Parser
    csv_content = b"id,topic,score\n1,Quantum Error Mitigation,0.95\n2,Zero Noise Extrapolation,0.92\n3,Surface Codes,0.89\n"
    pages_csv = parse_document(csv_content, "benchmark.csv", "text/csv")
    assert len(pages_csv) == 1
    assert "Quantum Error Mitigation" in pages_csv[0].text
    assert pages_csv[0].metadata["rows"] == 3
    assert pages_csv[0].metadata["columns"] == 3

    # 2. Test In-Memory PDF Parser (PyMuPDF)
    doc_pdf = fitz.open()
    page1 = doc_pdf.new_page()
    page1.insert_text((50, 72), "Page 1: Zero-noise extrapolation techniques.")
    page2 = doc_pdf.new_page()
    page2.insert_text((50, 72), "Page 2: Probabilistic error cancellation on superconducting qubits.")
    pdf_bytes = doc_pdf.write()
    doc_pdf.close()

    pages_pdf = parse_document(pdf_bytes, "paper.pdf", "application/pdf")
    assert len(pages_pdf) == 2
    assert "Page 1" in pages_pdf[0].text
    assert pages_pdf[0].page_number == 1
    assert "Page 2" in pages_pdf[1].text
    assert pages_pdf[1].page_number == 2

    # 3. Test In-Memory DOCX Parser
    doc_word = docx.Document()
    doc_word.add_heading("Quantum Telemetry Report", level=1)
    doc_word.add_paragraph("Empirical fidelity measurements over 127 physical qubits.")
    table = doc_word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Fidelity"
    table.cell(1, 1).text = "99.4%"
    docx_stream = io.BytesIO()
    doc_word.save(docx_stream)
    docx_bytes = docx_stream.getvalue()

    pages_docx = parse_document(docx_bytes, "report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert len(pages_docx) >= 1
    assert "Quantum Telemetry Report" in pages_docx[0].text
    assert "Fidelity" in pages_docx[0].text

    # 4. Strict Chunker Validation
    long_text = " ".join([f"Sentence {i} describing quantum error suppression factor {i * 0.1:.1f}." for i in range(1, 100)])
    pages_long = parse_document(long_text.encode("utf-8"), "large_corpus.txt", "text/plain")
    chunker = SemanticChunker(target_chunk_size=150, overlap_chars=30)
    chunks = chunker.chunk_pages(pages_long)
    assert len(chunks) >= 3

    # Check contiguous chunk indices and non-empty token counts
    for idx, chunk in enumerate(chunks):
        assert chunk.chunk_index == idx
        assert chunk.token_count > 0
        assert len(chunk.content) > 0
        assert chunk.metadata["source_filename"] == "large_corpus.txt"

    # Verify overlap exists between consecutive chunks
    overlap_found = False
    for i in range(len(chunks) - 1):
        words_curr = set(chunks[i].content.split()[-5:])
        words_next = set(chunks[i + 1].content.split()[:10])
        if words_curr.intersection(words_next):
            overlap_found = True
            break
    assert overlap_found, "Sliding overlap between consecutive chunks should share token n-grams"

    # 5. Strict Embedding Engine Mathematical Validation
    v1 = embedding_engine.embed_text("Quantum error mitigation in NISQ")
    v1_dup = embedding_engine.embed_text("Quantum error mitigation in NISQ")
    v2 = embedding_engine.embed_text("Quantum error extrapolation in NISQ devices")
    v3 = embedding_engine.embed_text("Culinary recipes for Italian pasta carbonara")

    # Dimensions
    assert len(v1) == 1536
    assert len(v2) == 1536
    assert len(v3) == 1536

    # Determinism: exact same vector for identical input
    assert v1 == v1_dup

    # Strict L2 Normalization check: ||v||_2 must be ~ 1.0 (unit vector)
    norm_v1 = math.sqrt(sum(x * x for x in v1))
    norm_v2 = math.sqrt(sum(x * x for x in v2))
    norm_v3 = math.sqrt(sum(x * x for x in v3))
    assert abs(norm_v1 - 1.0) < 1e-4, f"Vector norm {norm_v1} is not unit-length"
    assert abs(norm_v2 - 1.0) < 1e-4
    assert abs(norm_v3 - 1.0) < 1e-4

    # Cosine similarities: Related domain vs Unrelated domain
    sim_similar = sum(a * b for a, b in zip(v1, v2))
    sim_unrelated = sum(a * b for a, b in zip(v1, v3))
    assert sim_similar > sim_unrelated, f"Expected {sim_similar} > {sim_unrelated}"

    # Self-similarity must be ~ 1.0
    sim_self = sum(a * b for a, b in zip(v1, v1))
    assert abs(sim_self - 1.0) < 1e-4


@pytest.mark.asyncio
async def test_ingest_document_end_to_end(db_session, mock_cache_service):
    file_service = FileService(db_session, cache=mock_cache_service)
    ingestion_service = IngestionService(db_session)

    # Register user & project first
    user = User(email="analyst@nexus.org", password_hash="pw", name="Analyst")
    db_session.add(user)
    await db_session.flush()

    project = Project(name="Ingestion Test Project")
    db_session.add(project)
    await db_session.flush()

    member = ProjectMember(project_id=project.id, user_id=user.id, role="owner")
    db_session.add(member)
    await db_session.flush()

    # Register file
    file_dto = await file_service.register_file_upload(
        user.id,
        FileCreateDTO(
            project_id=project.id,
            filename="quantum_notes.txt",
            mime_type="text/plain",
            size_bytes=256,
            storage_key="uploads/quantum_notes.txt",
        ),
    )

    # Ingest document
    raw_doc = b"Quantum error mitigation reduces noise without requiring full fault-tolerant logical qubits.\n\nTechniques include Zero-Noise Extrapolation (ZNE) and Probabilistic Error Cancellation (PEC)."
    chunks = await ingestion_service.ingest_document(
        file_id=file_dto.id,
        content=raw_doc,
        filename="quantum_notes.txt",
        mime_type="text/plain",
    )

    assert len(chunks) >= 1
    assert chunks[0].file_id == file_dto.id
    assert chunks[0].token_count > 0
    assert len(chunks[0].content) > 10
    assert chunks[0].chunk_index == 0
